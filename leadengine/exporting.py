"""Lead exports: CSV, Excel (styled), Google Sheets (optional, service account)."""

from __future__ import annotations

import csv
import io
from typing import Any

from sqlalchemy.orm import Session

from leadengine.db.models import Business
from leadengine.db.repo import Repository

COLUMNS = [
    ("label", "Label"), ("opportunity", "Opportunity"), ("reason", "Why"), ("name", "Business"),
    ("status", "CRM status"), ("rating", "Rating"), ("reviews", "Reviews"), ("last_review", "Last review"),
    ("ads", "Google Ads"), ("lsa", "LSA"), ("site_score", "Website score"), ("site_grade", "Website grade"),
    ("site_flags", "Website flags"), ("site_reasons", "Website weaknesses"), ("email", "Best email"),
    ("email_conf", "Email confidence"), ("email_status", "Email verified"), ("owner", "Owner"),
    ("phone", "Phone"), ("website", "Website"), ("address", "Address"), ("city", "City"), ("state", "State"),
    ("zip", "ZIP"), ("categories", "Categories"), ("maps_url", "Google Maps"), ("place_id", "Place ID"),
]


def lead_rows(session: Session, businesses: list[Business]) -> list[dict[str, Any]]:
    repo = Repository(session)
    ids = [b.id for b in businesses]
    webs = repo.latest_enrichments(ids, "website")
    statuses = repo.lead_statuses(ids if len(ids) <= 900 else None)
    rows = []
    for b in businesses:
        web = webs.get(b.id)
        rows.append({
            "label": b.lead_label, "opportunity": b.opportunity_score, "reason": b.lead_reason, "name": b.name,
            "status": statuses.get(b.id, "New"), "rating": b.rating, "reviews": b.review_count,
            "last_review": b.last_review_at.date().isoformat() if b.last_review_at else None,
            "ads": b.ads_status, "lsa": "yes" if b.lsa else ("no" if b.lsa is False else None),
            "site_score": b.website_score, "site_grade": b.website_grade,
            "site_flags": ", ".join(b.website_flags or []),
            "site_reasons": "; ".join(((web.payload or {}).get("reasons") or [])[:4]) if web else None,
            "email": b.best_email, "email_conf": b.email_confidence, "email_status": b.email_status,
            "owner": b.owner_name, "phone": b.phone, "website": b.website, "address": b.address, "city": b.city,
            "state": b.state, "zip": b.zip_code, "categories": ", ".join(b.categories or []),
            "maps_url": b.google_maps_url, "place_id": b.place_id,
        })
    return rows


def to_csv(rows: list[dict[str, Any]]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([title for _, title in COLUMNS])
    for r in rows:
        writer.writerow(["" if r.get(k) is None else r.get(k) for k, _ in COLUMNS])
    return "﻿" + buf.getvalue()  # BOM so Excel opens UTF-8 correctly


def to_xlsx(rows: list[dict[str, Any]]) -> bytes:
    """Styled Excel file. Write-only mode keeps it fast for tens of thousands of rows."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Leads")
    widths = {"reason": 60, "name": 32, "site_reasons": 50, "address": 36, "website": 30, "email": 30,
              "maps_url": 30, "categories": 28}
    for i, (key, _) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(key, 14)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2937")
    header = []
    for _, title in COLUMNS:
        cell = WriteOnlyCell(ws, value=title)
        cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(vertical="center")
        header.append(cell)
    ws.append(header)
    fills = {k: PatternFill("solid", fgColor=v) for k, v in
             {"Hot": "FDE2E1", "Warm": "FEF3C7", "Cold": "E0F2FE", "Skip": "F3F4F6"}.items()}
    for r in rows:
        fill = fills.get(r.get("label"))
        if fill is None:
            ws.append([r.get(k) for k, _ in COLUMNS])
            continue
        line = []
        for k, _ in COLUMNS:
            cell = WriteOnlyCell(ws, value=r.get(k))
            cell.fill = fill
            line.append(cell)
        ws.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_google_sheet(rows: list[dict[str, Any]], *, credentials_file: str, spreadsheet: str,
                    worksheet: str = "Leads") -> str:
    """Write rows to a Google Sheet with a service account (``pip install gspread``).
    ``spreadsheet`` is a sheet URL or key; share the sheet with the service-account email first."""
    try:
        import gspread
    except ImportError as exc:
        raise RuntimeError("Google Sheets export needs: pip install gspread") from exc
    gc = gspread.service_account(filename=credentials_file)
    sh = gc.open_by_url(spreadsheet) if spreadsheet.startswith("http") else gc.open_by_key(spreadsheet)
    try:
        ws = sh.worksheet(worksheet)
        ws.clear()
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(title=worksheet, rows=len(rows) + 10, cols=len(COLUMNS))
    values = [[t for _, t in COLUMNS]] + [["" if r.get(k) is None else r.get(k) for k, _ in COLUMNS] for r in rows]
    ws.update(values, "A1")
    return sh.url
