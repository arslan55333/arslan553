import asyncio
import io
from datetime import timedelta

import pytest
from fastapi.testclient import TestClient

from leadengine import crm, jobs
from leadengine.db import Business, Job, Repository, utcnow
from leadengine.models import BusinessRecord, SearchQuery


def run(coro):
    return asyncio.run(coro)


def add(session, name, **kw):
    b = Repository(session).upsert_business(BusinessRecord(name=name, provider="t", place_id="P-" + name, **kw))
    session.commit()
    return b


# ── CRM ──────────────────────────────────────────────────────────────
def test_status_flow_and_never_contact_twice(session_factory):
    with session_factory() as s:
        a = add(s, "Acme", website="https://acme.com", phone="2145550001", rating=4.8, review_count=100)
        twin = add(s, "Acme Septic LLC", website="https://www.acme.com/septic", phone="2145559999")
        other = add(s, "Other", website="https://other.com", phone="2145550002")
        crm.set_status(s, a.id, "Preview Built")
        crm.set_status(s, a.id, "Emailed", "intro sent")
        s.commit()
        assert crm.current_status(s, a.id) == "Emailed" and s.get(Business, a.id).lead_label == "Skip"
        with pytest.raises(crm.AlreadyContacted, match="already contacted"):
            crm.set_status(s, a.id, "Emailed")
        with pytest.raises(crm.AlreadyContacted, match="looks like Acme"):
            crm.set_status(s, twin.id, "Emailed")                     # same domain -> same business
        crm.set_status(s, twin.id, "Emailed", force=True)              # explicit override allowed
        crm.set_status(s, other.id, "Emailed")
        crm.set_status(s, a.id, "Replied", "wants a call")
        crm.add_note(s, a.id, "call Tuesday")
        s.commit()
        events = crm.timeline(s, a.id)
        assert events[0].note == "call Tuesday" and events[1].status == "Replied"
        with pytest.raises(Exception):
            crm.set_status(s, a.id, "Bogus")


# ── job queue ────────────────────────────────────────────────────────
def test_job_runner_success_retry_resume_cancel(session_factory):
    calls = {"n": 0}

    async def flaky(ctx):
        for step in ctx.params["steps"]:
            if ctx.is_done(step):
                ctx.log(f"skip {step}")
                continue
            calls["n"] += 1
            if step == "b" and calls["n"] == 2:
                raise RuntimeError("network hiccup")
            ctx.step_done(step)
            ctx.log(f"did {step}")
        return {"ok": True}

    async def bad(ctx):
        raise ValueError("bad params")

    runner = jobs.JobRunner(session_factory, {"flaky": flaky, "bad": bad}, max_attempts=3)
    j1 = jobs.enqueue(session_factory, "flaky", {"steps": ["a", "b", "c"]})
    j2 = jobs.enqueue(session_factory, "bad", {})
    j3 = jobs.enqueue(session_factory, "flaky", {"steps": ["x"]})
    jobs.cancel(session_factory, j3)

    async def go():
        while await runner.run_once():
            pass
    run(go())
    with session_factory() as s:
        first, second, third = s.get(Job, j1), s.get(Job, j2), s.get(Job, j3)
        assert first.status == "done" and first.attempts == 2 and first.done_steps == ["a", "b", "c"]
        assert any("skip a" in line for line in first.progress)       # resumed without redoing step a
        assert any("will retry" in line for line in first.progress)
        assert second.status == "failed" and "bad params" in second.error   # bad input is not retried
        assert third.status == "cancelled"


def test_recover_stale_jobs(session_factory):
    j = jobs.enqueue(session_factory, "x", {})
    with session_factory() as s:
        job = s.get(Job, j)
        job.status, job.heartbeat_at = "running", utcnow() - timedelta(hours=1)
        s.commit()
    assert jobs.recover_stale(session_factory) == [j]
    with session_factory() as s:
        assert s.get(Job, j).status == "queued"


# ── web dashboard ────────────────────────────────────────────────────
@pytest.fixture
def client(settings, session_factory):
    from leadengine.ui.app import create_app

    app = create_app(settings, start_runner=False)
    with TestClient(app) as c:
        c.sf = app.state.sf
        yield c


def seed(sf):
    with sf() as s:
        repo = Repository(s)
        hot = add(s, "Hot Septic", website="https://hotseptic.com", phone="2145550001", rating=4.9, review_count=300,
                  zip_code="75201")
        hot.ads_status, hot.website_score, hot.best_email, hot.email_status = "Active", 12, "a@hotseptic.com", "valid"
        cold = add(s, "Cold Co", website="https://cold.com", phone="2145550002", rating=4.1, review_count=9,
                   zip_code="75201")
        cold.ads_status, cold.website_score = "None", 70
        repo.record_search(SearchQuery("septic", "75201"), "t", [(hot, 1), (cold, 2)], exhausted=True, api_calls=1)
        repo.set_enrichment(hot.id, "website", {"score": 12, "reasons": ["copyright 2011 (~15 years without updates)"],
                                                "flags": [], "signals": {}})
        s.commit()
        return hot.id, cold.id


def test_pages_filters_and_exports(client):
    from leadengine.service import LeadService

    hot_id, cold_id = seed(client.sf)
    LeadService(client.app.state.settings, client.sf, None).rescore()
    home = client.get("/")
    assert home.status_code == 200 and "Hot Septic" in home.text
    page = client.get("/leads?label=Hot").text
    assert "Hot Septic" in page and "Cold Co" not in page
    no_ads = client.get("/leads?ads=None").text
    assert "Cold Co" in no_ads and "Hot Septic" not in no_ads
    assert "Hot Septic" in client.get("/leads?q=hotseptic").text
    assert "copyright 2011" in client.get(f"/leads/{hot_id}").text
    csv_text = client.get("/export.csv?label=Hot").text
    assert csv_text.startswith("﻿Label,Opportunity,Why,Business")
    assert "Hot Septic" in csv_text and "Cold Co" not in csv_text
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(client.get("/export.xlsx").content))["Leads"]
    assert ws["D1"].value == "Business" and {ws["D2"].value, ws["D3"].value} == {"Hot Septic", "Cold Co"}
    assert client.get("/credits").status_code == 200 and client.get("/jobs").status_code == 200


def test_status_change_and_duplicate_guard_in_ui(client):
    hot_id, _ = seed(client.sf)
    r = client.post(f"/leads/{hot_id}/status", data={"status": "Emailed", "note": "sent"}, follow_redirects=False)
    assert r.status_code == 303
    again = client.post(f"/leads/{hot_id}/status", data={"status": "Emailed"}, follow_redirects=True)
    assert "already contacted" in again.text
    with client.sf() as s:
        assert crm.current_status(s, hot_id) == "Emailed"


def test_run_form_creates_job_and_log_endpoint(client):
    r = client.post("/run", data={"keyword": "septic", "zips": "75201, 75204 bad 7520", "emails": "true"},
                    follow_redirects=False)
    assert r.status_code == 303
    job_id = int(r.headers["location"].rsplit("/", 1)[1])
    with client.sf() as s:
        job = s.get(Job, job_id)
        assert job.params["zips"] == ["75201", "75204"] and job.params["options"]["emails"] is True
        assert job.params["options"]["ads"] is False
    log = client.get(f"/jobs/{job_id}/log").text
    assert "every 2s" in log and "0/2 ZIPs" in log            # keeps polling while queued
    assert client.post("/run", data={"keyword": "x", "zips": "abc"}).status_code == 400


def test_media_is_confined_to_data_dir(client, settings):
    shots = settings.root / "data" / "screenshots"
    shots.mkdir(parents=True, exist_ok=True)
    (shots / "1.jpg").write_bytes(b"\xff\xd8\xffjpeg")
    assert client.get("/media/screenshots/1.jpg").content.startswith(b"\xff\xd8")
    assert client.get("/media/../config.toml").status_code == 404
    assert client.get("/media/%2e%2e/config.toml").status_code == 404


def test_discover_job_handler_skips_finished_zips(settings, session_factory, monkeypatch):
    from leadengine import workers
    from leadengine.service import DiscoverOutcome, LeadService

    seen = []

    async def fake_discover(self, keyword, zip_code, **kw):
        seen.append((zip_code, {k: v for k, v in kw.items() if k != "on_progress"}))
        return DiscoverOutcome(1, "playwright", zip_code, False, [], sponsored=1)

    monkeypatch.setattr(LeadService, "discover", fake_discover)
    handler = workers.make_handlers(settings, session_factory)["discover"]
    params = {"keyword": "septic", "zips": ["75201", "75204"], "options": {"emails": False, "bogus": 1}}
    job_id = jobs.enqueue(session_factory, "discover", params)
    ctx = jobs.JobContext(session_factory, job_id, params)
    ctx.step_done("75201")                                      # pretend the first ZIP finished before a crash
    out = run(handler(ctx))
    assert [z for z, _ in seen] == ["75204"] and seen[0][1] == {"emails": False}
    assert out == [{"zip": "75204", "businesses": 0, "hot": 0, "warm": 0, "sponsored": 1, "from_cache": False}]
